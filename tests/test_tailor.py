"""Tailor tests: fact checker, rendering limits, refusal and retry behaviour. The model is faked."""
import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml

from config import OPUS, SONNET
from tailor import layout, render, tailor as t
from tailor.facts import Facts
from tailor.factcheck import Context, check_text
from tailor.plan import TailorPlan

MASTER = {
    "identity": {"name": "Asha Rao", "email": "asha@example.com", "phone": "+91-9000000000",
                 "linkedin": "linkedin.com/in/asha", "portfolio": "asha.dev"},
    "headline_variants": {"default": "Product Manager with 7+ years building B2B SaaS and consumer apps",
                          "saas": "Product Manager who scaled a SaaS product to ₹25Cr ARR"},
    "summary_facts": ["7 years of product experience", "Scaled a SaaS product to ₹25Cr ARR before an acquisition",
                      "Comfortable as the first or only PM"],
    "skills": {"product": ["Roadmapping", "A/B testing"], "tools": ["Salesforce", "Figma"]},
    "experience": [
        {"id": "r1", "company": "Beta Pvt. Ltd.", "title": "Associate Vice President - Product",
         "location": "Hyderabad", "start": "August 2024", "end": "July 2026", "context": "D2C brand",
         "achievements": [
             {"id": "a1", "tags": ["growth", "subscription"], "metrics": ["+85% session time"],
              "text": "Re-architected the subscription app, driving +85% session time and -63% time to subscribe."},
             {"id": "a2", "tags": ["crm", "salesforce"], "metrics": ["4 teams"],
              "text": "Shipped a Salesforce CRM unifying WhatsApp and agent workflows across 4 teams."},
             {"id": "a3", "tags": ["logistics"], "metrics": ["₹90L saved"],
              "text": "Cut ₹90L in delivery cost by geofencing dark stores."}]},
        {"id": "r2", "company": "Gamma Technologies Pvt. Ltd.", "title": "Product Manager", "location": "Bengaluru",
         "start": "October 2021", "end": "June 2024", "context": "B2B SaaS that scaled to ₹25Cr ARR",
         "achievements": [
             {"id": "b1", "tags": ["saas", "payments"], "metrics": ["5 payment methods"],
              "text": "Built the Payment module from zero with 5 payment methods."},
             {"id": "b2", "tags": ["process"], "metrics": ["+33% delivery speed"],
              "text": "Introduced Agile practices: +33% delivery speed."}]}],
    "education": {"institution": "Nitte Institute", "degree": "Bachelor of Engineering", "location": "Bengaluru",
                  "start": "August 2015", "end": "May 2019"},
    "certifications": ["Scrum Alliance: CSPO Certification"],
    "equivalents": [{"inr": "₹25Cr ARR", "usd": "$3M ARR"}],
}
FILLER = "I have shipped product work like this before and I would like to help your team do the same with care"


def n_words(n):
    ws = (FILLER + " ").split() * 20
    return " ".join(ws[:n]) + "."


def good_plan(**over):
    plan = {"headline_key": "saas", "summary": ["7 years of product experience.", "Comfortable as the first or only PM."],
            "experience": [
                {"id": "r1", "bullets": [
                    {"achievement_id": "a2", "text": MASTER["experience"][0]["achievements"][1]["text"], "reword_why": ""},
                    {"achievement_id": "a1", "text": "Re-architected the subscription app and grew session time, +85% session time, -63% time to subscribe.",
                     "reword_why": "mirrors the JD's growth wording"}]},
                {"id": "r2", "bullets": [
                    {"achievement_id": "b1", "text": MASTER["experience"][1]["achievements"][0]["text"], "reword_why": ""}]}],
            "skills": {"product": ["Roadmapping"], "tools": ["Salesforce"]},
            "certifications": ["Scrum Alliance: CSPO Certification"],
            "jd_keywords": ["roadmapping", "salesforce", "kubernetes"],
            "selection_rationale": "CRM and growth work match the JD.",
            "cover_note": n_words(135), "outreach": n_words(40)}
    plan.update(over)
    return plan


class FakeClient:
    def __init__(self, replies):
        self.replies, self.calls = list(replies), []
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kw):
        self.calls.append(kw)
        text = self.replies.pop(0)
        return SimpleNamespace(content=[SimpleNamespace(text=text if isinstance(text, str) else json.dumps(text))],
                               usage=SimpleNamespace(input_tokens=3000, output_tokens=1500))


needs_font = pytest.mark.skipif(not any(Path(a).exists() for a, _ in layout.FONT_CANDIDATES),
                                reason="needs a TTF font with the rupee sign")


@pytest.fixture
def env(tmp_path):
    data = tmp_path / "data" / "clean"
    data.mkdir(parents=True)

    def row(jid, score=72, location="Bangalore, Karnataka", partial=False, company="Acme Corp"):
        return {"id": jid, "source": "greenhouse", "title": "Senior Product Manager", "company": company,
                "location": location, "url": "https://x", "description": "Own the roadmap. Salesforce and Jira a plus.",
                "jd_partial": partial, "also_listed_on": [],
                "screen": {"status": "scored", "stage2": {"score": score, "verdict": "apply"}}}
    rows = [row("in1"), row("us1", location="Remote - US"), row("top", score=90), row("part", partial=True)]
    (data / "screened_jobs.json").write_text(json.dumps(rows))
    (tmp_path / "master.json").write_text(json.dumps(MASTER))
    (tmp_path / "profile.yaml").write_text(yaml.safe_dump({"scoring": {"thresholds": {"top_pick_min_score": 85}}}))
    return SimpleNamespace(root=tmp_path, out=tmp_path / "out" / "resumes")


def run(env, job_id="in1", client=None, **kw):
    return t.tailor_job(job_id, data_dir=env.root / "data", master_path=env.root / "master.json",
                        profile_path=env.root / "profile.yaml", out_root=env.out, client=client,
                        today=date(2026, 10, 16), **kw)


# ---------- fact checker ----------

FACTS = Facts.from_master(MASTER)
CTX = Context(company="Acme Corp", title="Senior Product Manager", jd_text="Own the roadmap. Jira and Kubernetes a plus.")


def violations(text):
    return check_text({"t": text}, FACTS, CTX)


def test_clean_text_passes():
    assert violations("Re-architected the subscription app, driving +85% session time. Cut ₹90L in delivery cost. "
                      "Product Manager at Gamma Technologies Pvt. Ltd. from October 2021 to June 2024. Salesforce CRM.") == []


def test_invented_number_is_caught():
    v = violations("Grew session time by +95% across 4 teams.")
    assert any("95%" in x for x in v)
    assert any("₹90Cr" in x for x in violations("Cut ₹90Cr in delivery cost."))   # unit swapped from ₹90L
    assert any("nine" in x for x in violations("Led nine designers."))        # number words count


def test_invented_tool_is_caught_including_jd_tools():
    assert any("jira" in x for x in violations("Managed the roadmap in Jira."))
    assert any("kubernetes" in x for x in violations("Shipped on Kubernetes."))
    assert violations("Managed work in Salesforce and Figma.") == []              # in skills.tools


def test_changed_title_or_date_is_caught():
    assert any("Group Product Manager" in x for x in violations("I was a Group Product Manager at Gamma."))
    assert any("Senior Director of Product" in x for x in violations("As Senior Director of Product I led teams."))
    assert any("March 2020" in x for x in violations("Worked there from March 2020."))
    assert any("2022" in x for x in violations("Product Manager from October 2022."))
    assert violations("Associate Vice President - Product, August 2024 to July 2026.") == []


def test_invented_company_is_caught():
    assert any("Zoho" in x for x in violations("Earlier I worked at Zoho Pvt Ltd on payments."))
    assert any("Flipkart" in x for x in violations("I joined Flipkart to run checkout."))


def test_inr_and_usd_equivalent_pair_is_one_fact():
    assert violations("Scaled the product to ₹25Cr ARR.") == []
    assert violations("Scaled the product to $3M ARR.") == []
    assert any("$4M" in x for x in violations("Scaled the product to $4M ARR."))


def test_em_dash_is_a_violation():
    assert any("em dash" in x for x in violations("Built payments — fast."))


def test_localization_picks_currency_by_job_location():
    assert t.is_india_job("Bangalore, Karnataka") and t.is_india_job("Remote - India")
    assert not t.is_india_job("Remote - US") and not t.is_india_job("N/A")
    assert "$3M ARR" in FACTS.localize("scaled to ₹25Cr ARR", india=False)
    assert "₹25Cr ARR" in FACTS.localize("scaled to $3M ARR", india=True)


def test_prompt_shows_one_currency_only(env):
    kw = dict(data_dir=env.root / "data", master_path=env.root / "master.json", profile_path=env.root / "profile.yaml")
    us = t.prepare("us1", **kw).user
    ind = t.prepare("in1", **kw).user
    assert "$3M ARR" in us and "₹25Cr" not in us
    assert "₹25Cr ARR" in ind and "$3M" not in ind


# ---------- plan structure ----------

def plan_violations(plan_dict):
    return t.structure_violations(TailorPlan.model_validate(plan_dict), FACTS)


def test_structure_rules():
    assert plan_violations(good_plan()) == []
    assert any("headline_key" in v for v in plan_violations(good_plan(headline_key="nope")))
    assert any("skill 'Kubernetes'" in v for v in plan_violations(good_plan(skills={"tools": ["Kubernetes"]})))
    assert any("cover_note" in v for v in plan_violations(good_plan(cover_note=n_words(90))))
    assert any("outreach" in v for v in plan_violations(good_plan(outreach=n_words(70))))
    bad = good_plan()
    bad["experience"][0]["bullets"][1]["text"] = "Re-architected the app, +95% session time."
    assert any("bullet a1" in v and "95%" in v for v in plan_violations(bad))
    bad = good_plan()
    bad["experience"][0]["bullets"][0]["text"] = "Shipped a Salesforce CRM across 4 teams, reworded."
    assert any("reword_why is empty" in v for v in plan_violations(bad))
    bad = good_plan()
    bad["experience"][0]["bullets"].append({"achievement_id": "b2", "text": "x", "reword_why": "y"})
    assert any("not an achievement of role r1" in v for v in plan_violations(bad))


def test_bullet_cannot_borrow_a_number_from_another_achievement():
    bad = good_plan()
    bad["experience"][0]["bullets"][1]["text"] = "Re-architected the subscription app, +85% session time across 4 teams."
    assert any("bullet a1" in v and "'4'" in v for v in plan_violations(bad))


# ---------- end to end (fake model) ----------

@needs_font
def test_happy_path_writes_the_package(env):
    client = FakeClient([good_plan()])
    res = run(env, client=client)
    assert res.folder.name == "Acme-Corp_Senior-Product-Manager_20261016" and res.calls == 1
    assert res.model == SONNET
    assert sorted(p.name for p in res.folder.iterdir()) == sorted(
        ["resume.docx", "resume.pdf", "cover_note.md", "outreach.md", "changes.md", "jd_keywords.md"])
    pdf = (res.folder / "resume.pdf").read_bytes()
    assert render.pdf_page_count(pdf) <= 2
    text = render.pdf_text(pdf)
    assert "Gamma Technologies Pvt. Ltd." in text and "₹25Cr ARR" in text     # real text, rupee sign intact
    assert "—" not in text
    words = len((res.folder / "cover_note.md").read_text().split())
    assert 120 <= words <= 150 and len((res.folder / "outreach.md").read_text().split()) < 70
    changes = (res.folder / "changes.md").read_text()
    assert "Reworded a1" in changes and "Before:" in changes and "reordered" in changes
    kw = (res.folder / "jd_keywords.md").read_text()
    assert "salesforce: backed" in kw and "kubernetes: NOT backed" in kw
    from docx import Document
    doc = Document(str(res.folder / "resume.docx"))
    assert doc.tables == [] and len(doc.inline_shapes) == 0
    assert "Asha Rao" in "\n".join(p.text for p in doc.paragraphs)


@needs_font
def test_us_job_shows_usd_and_top_score_uses_opus(env):
    res = run(env, "us1", FakeClient([good_plan()]))
    assert "$3M ARR" in render.pdf_text((res.folder / "resume.pdf").read_bytes())
    assert run(env, "top", FakeClient([good_plan()])).model == OPUS


def test_partial_jd_refused_until_full_jd_pasted(env):
    client = FakeClient([good_plan()])
    with pytest.raises(t.PartialJDError):
        run(env, "part", client)
    assert client.calls == [] and not env.out.exists()


@needs_font
def test_pasted_jd_is_stored_only_in_the_output_folder(env):
    res = run(env, "part", FakeClient([good_plan()]), jd_override="Full JD text. Own the roadmap.")
    assert (res.folder / "jd_pasted.txt").read_text() == "Full JD text. Own the roadmap."
    assert "Full JD text" not in json.dumps(json.loads((env.root / "data/clean/screened_jobs.json").read_text()))


def test_retry_then_fail_writes_no_files(env):
    bad = good_plan()
    bad["experience"][0]["bullets"][1]["text"] = "Re-architected the app, +95% session time."
    client = FakeClient([bad, bad])
    with pytest.raises(t.FactCheckFailed) as e:
        run(env, client=client)
    assert len(client.calls) == 2 and any("95%" in v for v in e.value.violations)
    assert not env.out.exists() or list(env.out.iterdir()) == []


@needs_font
def test_retry_lists_violations_and_can_succeed(env):
    bad = good_plan(cover_note=n_words(90))
    client = FakeClient([bad, good_plan()])
    res = run(env, client=client)
    assert res.calls == 2 and res.violations_fixed
    retry_msg = client.calls[1]["messages"][-1]["content"]
    assert "cover_note is 90 words" in retry_msg and client.calls[1]["messages"][1]["role"] == "assistant"


def test_invalid_json_gets_one_retry(env):
    client = FakeClient(["not json at all", "still not json"])
    with pytest.raises(t.FactCheckFailed):
        run(env, client=client)
    assert len(client.calls) == 2


def test_cost_guard_blocks_before_any_call(env):
    client = FakeClient([good_plan()])
    with pytest.raises(t.CostGuardError):
        run(env, client=client, max_cost=0.0001)
    assert client.calls == []


# ---------- 2 page limit ----------

@needs_font
def test_pdf_is_verified_to_two_pages_and_long_content_is_trimmed(env):
    big = json.loads(json.dumps(MASTER))
    long_text = ("Built a platform feature with careful discovery and delivery across several teams. " * 6).strip()
    big["experience"][0]["achievements"] = [{"id": f"x{i}", "tags": [], "metrics": [], "text": long_text} for i in range(40)]
    plan = TailorPlan.model_validate(good_plan(experience=[
        {"id": "r1", "bullets": [{"achievement_id": f"x{i}", "text": long_text, "reword_why": ""} for i in range(40)]},
        {"id": "r2", "bullets": [{"achievement_id": "b1", "text": MASTER["experience"][1]["achievements"][0]["text"],
                                  "reword_why": ""}]}]))
    master_path = env.root / "big.json"
    master_path.write_text(json.dumps(big))
    prep = t.prepare("in1", data_dir=env.root / "data", master_path=master_path, profile_path=env.root / "profile.yaml")
    pdf, size, pages, dropped, _ = t.trim_to_fit(plan, prep)
    assert pages <= 2 and render.pdf_page_count(pdf) <= 2
    assert len(dropped) > 0                       # bullets were dropped, and the list says which
    name, contact, sections = t.build_sections(plan, prep)
    assert render.fit_pdf(name, contact, sections)[2] <= 2
