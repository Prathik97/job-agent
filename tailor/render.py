"""Render the resume as DOCX and PDF. Single column, standard headings, real text, no tables
or images. The PDF is measured, and the page limit is enforced on the generated PDF."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import List, Optional, Tuple

from tailor import layout

Block = Tuple[str, ...]   # ("para", text) | ("role", company, title, meta) | ("bullet", text)

# Metrics shown in bold, as in the reference resume: "+85%", "2x", "-63%", "₹18.5L", "₹25Cr", "500+", "~44%".
_METRIC = re.compile(r"[-+~]?(?:₹\s?)?\d[\d,.]*(?:\s?(?:%|x\b|Cr\b|L\b|\+)|(?<=₹\d)|(?<=₹\s\d))"
                     r"|[-+~]?₹\s?\d[\d,.]*")


class RenderError(Exception):
    pass


def pdf_page_count(data: bytes) -> int:
    from pypdf import PdfReader
    return len(PdfReader(io.BytesIO(data)).pages)


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)


def split_metrics(text: str) -> List[Tuple[str, bool]]:
    """Split text into (segment, is_metric) so numbers can be bolded. Wording is unchanged."""
    parts, pos = [], 0
    for m in _METRIC.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(0), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts


def _register_font() -> Tuple[str, str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold, italic in layout.FONT_CANDIDATES:
        if all(Path(f).exists() for f in (regular, bold, italic)):
            pdfmetrics.registerFont(TTFont(layout.FONT_NAME, regular))
            pdfmetrics.registerFont(TTFont(layout.FONT_NAME + "-Bold", bold))
            pdfmetrics.registerFont(TTFont(layout.FONT_NAME + "-Italic", italic))
            pdfmetrics.registerFontFamily(layout.FONT_NAME, normal=layout.FONT_NAME,
                                          bold=layout.FONT_NAME + "-Bold", italic=layout.FONT_NAME + "-Italic",
                                          boldItalic=layout.FONT_NAME + "-Bold")   # makes <b> use real bold
            return layout.FONT_NAME, layout.FONT_NAME + "-Bold", layout.FONT_NAME + "-Italic"
    raise RenderError("no font with the rupee sign found; add one to tailor/layout.py FONT_CANDIDATES")


def build_pdf(name: str, contact: str, sections: List[Tuple[str, List[Block]]], size: float) -> bytes:
    """`contact` is the headline line, a newline, then the contact line (both centred)."""
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (HRFlowable, KeepTogether, ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer)
    from xml.sax.saxutils import escape

    regular, bold, italic = _register_font()

    def rich(text: str) -> str:
        return "".join(f"<b>{escape(t)}</b>" if is_metric else escape(t) for t, is_metric in split_metrics(text))

    base = ParagraphStyle("base", fontName=regular, fontSize=size, leading=size * 1.28, alignment=TA_LEFT)
    centre = ParagraphStyle("centre", parent=base, alignment=TA_CENTER)
    title = ParagraphStyle("name", parent=centre, fontName=bold, fontSize=size + 6, leading=(size + 6) * 1.2)
    head = ParagraphStyle("head", parent=base, fontName=bold, fontSize=size, spaceBefore=size * 0.8,
                          spaceAfter=0, keepWithNext=1)
    company = ParagraphStyle("company", parent=base, fontName=bold, spaceBefore=size * 0.45, keepWithNext=1)
    meta = ParagraphStyle("meta", parent=base, fontName=italic, keepWithNext=1)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=layout.MARGIN_INCH * inch, rightMargin=layout.MARGIN_INCH * inch,
                            topMargin=layout.MARGIN_INCH * inch, bottomMargin=layout.MARGIN_INCH * inch,
                            title=f"{name} resume", author=name)
    story = [Paragraph(escape(name.upper()), title)]
    story += [Paragraph(escape(line), centre) for line in contact.split("\n")]
    for heading, blocks in sections:
        start = len(story)
        rule = HRFlowable(width="100%", thickness=0.8, spaceBefore=1, spaceAfter=size * 0.3)
        rule.keepWithNext = 1
        story += [Paragraph(escape(heading), head), rule]
        bullets: List = []

        def flush():
            if bullets:
                story.append(ListFlowable(list(bullets), bulletType="bullet", start="•", leftIndent=14,
                                          bulletFontName=regular, bulletFontSize=size))
                bullets.clear()
        for b in blocks:
            if b[0] == "bullet":
                bullets.append(ListItem(Paragraph(rich(b[1]), base), leftIndent=14, spaceAfter=size * 0.2))
                continue
            flush()
            if b[0] == "role":
                story += [Paragraph(escape(b[1]), company), Paragraph(escape(b[2]), base.clone("t", keepWithNext=1)),
                          Paragraph(escape(b[3]), meta)]
            else:
                story.append(Paragraph(rich(b[1]), base))
        flush()
        if len(blocks) <= 4:   # short sections stay on one page
            story[start:] = [KeepTogether(story[start:])]
    story.append(Spacer(1, 1))
    doc.build(story)
    return buf.getvalue()


def _bottom_border(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    border = OxmlElement("w:pBdr")
    line = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        line.set(qn(f"w:{key}"), value)
    border.append(line)
    paragraph._p.get_or_add_pPr().append(border)


def build_docx(name: str, contact: str, sections: List[Tuple[str, List[Block]]], size: float) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    def add_rich(paragraph, text: str) -> None:
        for chunk, is_metric in split_metrics(text):
            paragraph.add_run(chunk).bold = is_metric

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(layout.MARGIN_INCH)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = layout.FONT_NAME, Pt(size)
    normal.paragraph_format.space_after = Pt(1)
    for style in ("Heading 1", "List Bullet", "Title"):
        doc.styles[style].font.name = layout.FONT_NAME
    doc.styles["Heading 1"].font.size = Pt(size)
    doc.styles["Title"].font.size = Pt(size + 6)
    title = doc.add_paragraph(name.upper(), style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in contact.split("\n"):
        doc.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, blocks in sections:
        _bottom_border(doc.add_paragraph(heading, style="Heading 1"))
        for b in blocks:
            if b[0] == "bullet":
                add_rich(doc.add_paragraph(style="List Bullet"), b[1])
            elif b[0] == "role":
                doc.add_paragraph().add_run(b[1]).bold = True
                doc.add_paragraph(b[2])
                doc.add_paragraph().add_run(b[3]).italic = True
            else:
                add_rich(doc.add_paragraph(), b[1])
    doc.core_properties.author = name
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def fit_pdf(name: str, contact: str, sections: List[Tuple[str, List[Block]]]) -> Tuple[bytes, float, int]:
    """Smallest-change layout that fits: try each font size; returns (pdf, size, pages)."""
    pages = 0
    for size in layout.FONT_SIZES:
        pdf = build_pdf(name, contact, sections, size)
        pages = pdf_page_count(pdf)
        if pages <= layout.MAX_PAGES:
            return pdf, size, pages
    return pdf, layout.FONT_SIZES[-1], pages
